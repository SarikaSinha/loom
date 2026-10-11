import asyncio
import json
import os
import re
import runpy
import sqlite3
import threading
import time
from pathlib import Path

import numpy as np
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from loom import Agent, DocumentStore
from loom.config import Config, StoreConfig
from loom.contracts import DocumentAnswer, ToolCall, ToolFailure, canonical
from loom.providers.base import ModelResponse
from loom.security.crypto import Crypto
from loom.testing import ScriptedModel


class FakeEncoder:
    def __init__(self):
        self.batches = []

    def offsets(self, text):
        return [match.span() for match in re.finditer(r"\S+", text)]

    def embed(self, texts, *, query=False):
        if not query:
            self.batches.append(tuple(texts))
        vectors = np.zeros((len(texts), 384), dtype=np.float32)
        vectors[:, 0] = 1
        return vectors


def test_three_format_index_and_grounded_answer(tmp_path, encrypted_config, monkeypatch, capsys):
    from loom.documents import embeddings

    monkeypatch.setattr(embeddings, "encoder", lambda config: FakeEncoder())
    source = tmp_path / "product-docs"
    source.mkdir()
    (source / "limits.md").write_text("# Rate limits\n\nThe API allows 500 requests/minute.\n")
    document = Document()
    document.add_heading("FAQ", 1)
    document.add_paragraph("The API supports pagination.")
    document.save(str(source / "faq.docx"))
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (API authentication uses an API key.) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.write(source / "api-guide.pdf")
    example = Path(__file__).resolve().parents[2] / "sdd/query-agent-example"
    config_path = tmp_path / "loom.toml"
    config_path.write_text((example / "loom.toml").read_text())
    monkeypatch.setenv("LOOM_ENCRYPTION_KEY", os.environ["LOOM_TEST_KEY"])
    config = Config.load(config_path)
    store = DocumentStore(path=source).bind(config)
    report = store.build()
    assert report["indexed"] == 3
    results = store.search("API rate limits")
    assert results.snapshot_id == report["snapshot_id"]
    assert {chunk.location.kind for chunk in results.chunks} == {"pdf", "markdown", "docx"}
    for path in (store.index_path / "snapshots" / results.snapshot_id).iterdir():
        assert b"500 requests" not in path.read_bytes()

    def respond(request):
        observations = [message.observation for message in request.messages if message.observation]
        if not observations:
            return ModelResponse(
                finish_reason="tool_calls",
                tool_calls=(
                    ToolCall(
                        call_id="search",
                        provider_call_id="search",
                        name="search_product_docs",
                        arguments={"query": "API rate limits"},
                    ),
                ),
            )
        result = observations[-1].output
        chunk = next(chunk for chunk in result["chunks"] if "500 requests" in chunk["text"])
        return ModelResponse(
            output={
                "status": "answered",
                "answer": "The API allows 500 requests/minute.",
                "sources": [
                    {
                        "store": result["store"],
                        "snapshot_id": result["snapshot_id"],
                        "chunk_id": chunk["chunk_id"],
                        "path": chunk["path"],
                        "location": chunk["location"],
                        "excerpt": "The API\tallows 500 requests/minute.  ",
                    }
                ],
            }
        )

    agent = Agent(
        name="product_expert",
        instructions="Answer from documents and cite evidence.",
        model=ScriptedModel(respond),
        data=[store],
        config=config,
        output_schema="document_qa",
    )
    result = agent.run(input="What are the API rate limits?")
    assert result.status == "completed"
    assert result.output["sources"][0]["path"] == "limits.md"
    assert result.output["sources"][0]["excerpt"] == "The API allows 500 requests/minute."
    assert result.state.provenance[0].store == "product_docs"

    script = example / "developer-code.py"
    assert (
        DocumentAnswer.model_validate_json((example / "output.json").read_bytes()).status
        == "abstained"
    )
    import loom.providers.factory

    monkeypatch.setattr(
        loom.providers.factory, "create_adapter", lambda profile, config: ScriptedModel(respond)
    )
    monkeypatch.chdir(tmp_path)
    runpy.run_path(str(script), run_name="__main__")
    printed = json.loads(capsys.readouterr().out)
    assert printed["status"] == "completed"
    assert "500 requests/minute" in printed["output"]["answer"]
    assert "confidence" not in printed["output"]


def test_search_requires_explicit_binding(tmp_path):
    with pytest.raises(Exception, match="store binding"):
        DocumentStore(path=tmp_path).search("query")


@pytest.fixture
def document_store(tmp_path, encrypted_config, monkeypatch):
    from loom.documents import embeddings

    encoder = FakeEncoder()
    monkeypatch.setattr(embeddings, "encoder", lambda config: encoder)
    source = tmp_path / "docs"
    source.mkdir()
    config = encrypted_config.model_copy(
        update={
            "stores": {
                "docs": StoreConfig(path=str(source), index_path=str(tmp_path / "index")),
            }
        }
    )
    return DocumentStore(path=source).bind(config), source, encoder


def snapshot_rows(store, snapshot):
    _, image = Crypto(store.config).read(
        store.index_path / "snapshots" / snapshot / "index.sqlite.enc",
        "index_sqlite",
        512 * 1024**2,
    )
    database = sqlite3.connect(":memory:")
    try:
        database.deserialize(image)
        return database.execute(
            "SELECT chunk_id,source_path,text,embedding FROM chunks ORDER BY chunk_id"
        ).fetchall()
    finally:
        database.close()


def test_refresh_reuses_vectors_and_preserves_historical_snapshots(document_store):
    store, source, encoder = document_store
    (source / "one.md").write_text("Original first document.")
    (source / "two.md").write_text("Original second document.")
    first = store.build()
    original = snapshot_rows(store, first["snapshot_id"])
    encoder.batches.clear()
    second = store.build(refresh=True)
    assert second["snapshot_id"] != first["snapshot_id"]
    assert second["reused_vectors"] == second["chunk_count"] == 2
    assert encoder.batches == []
    assert snapshot_rows(store, second["snapshot_id"]) == original
    assert second["retained_bytes"] == first["snapshot_bytes"] + second["snapshot_bytes"]
    assert (
        Crypto(store.config)
        .read(store.index_path / "CURRENT.enc", "index_current", 65536)[0]
        .revision
        == 2
    )

    (source / "one.md").write_text("Changed first document.")
    (source / "two.md").unlink()
    third = store.build(refresh=True)
    assert encoder.batches == [("Changed first document.",)]
    assert third["reused_vectors"] == 0
    assert third["chunk_count"] == 1
    assert snapshot_rows(store, third["snapshot_id"])[0][2] == "Changed first document."
    assert snapshot_rows(store, first["snapshot_id"]) == original


def test_refresh_drops_newly_failed_files_and_returns_full_report(document_store):
    store, source, _ = document_store
    (source / "good.md").write_text("Retained content.")
    (source / "broken.md").write_text("Content that will fail on refresh.")
    first = store.build()
    (source / "broken.md").write_bytes(b"\xff")
    (source / "bad.pdf").write_bytes(b"not a PDF")
    (source / "bad.docx").write_bytes(b"not a ZIP")
    (source / "ignored.txt").write_text("Unsupported content.")
    report = store.build(refresh=True)
    assert report["indexed"] == 1
    assert report["excluded"] == 4
    assert report["reused_vectors"] == 1
    assert {row[1] for row in snapshot_rows(store, report["snapshot_id"])} == {"good.md"}
    assert len(snapshot_rows(store, first["snapshot_id"])) == 2
    reasons = {entry["path"]: entry.get("reason") for entry in report["files"]}
    assert reasons == {
        "bad.docx": "docx_parse_failed",
        "bad.pdf": "pdf_parse_failed",
        "broken.md": "parse_failed",
        "good.md": None,
        "ignored.txt": "unsupported_format",
    }
    _, persisted = Crypto(store.config).read_json(
        store.index_path / "snapshots" / report["snapshot_id"] / "ingestion-report.json.enc",
        "index_report",
    )
    assert persisted == report["files"]
    found = store.search("Retained")
    assert found.diagnostics[0].code == "index_coverage_incomplete"
    assert found.diagnostics[0].details["excluded"] == 4
    assert "bad.pdf" not in canonical(found.diagnostics[0]).decode()


def test_embedding_batches_and_pipeline_invalidation(document_store, monkeypatch):
    from loom.documents import index as documents

    store, source, encoder = document_store
    (source / "many.md").write_text("\n\n".join(f"Paragraph {i}." for i in range(67)))
    report = store.build()
    assert report["chunk_count"] == 67
    assert [len(batch) for batch in encoder.batches] == [32, 32, 3]
    encoder.batches.clear()
    monkeypatch.setattr(documents, "PIPELINE", "different_pipeline")
    refreshed = store.build(refresh=True)
    assert refreshed["reused_vectors"] == 0
    assert [len(batch) for batch in encoder.batches] == [32, 32, 3]


def test_empty_and_all_failed_refreshes(document_store):
    store, source, _ = document_store
    empty = store.build()
    assert empty["chunk_count"] == 0
    assert store.search("anything").chunks == ()
    pointer = (store.index_path / "CURRENT.enc").read_bytes()
    (source / "unsupported.txt").write_text("Not supported")
    with pytest.raises(ToolFailure) as failure:
        store.build(refresh=True)
    assert failure.value.code == "no_indexable_documents"
    assert (store.index_path / "CURRENT.enc").read_bytes() == pointer
    (source / "unsupported.txt").unlink()
    refreshed = store.build(refresh=True)
    assert refreshed["snapshot_id"] != empty["snapshot_id"]
    assert refreshed["chunk_count"] == 0


def test_failed_publication_preserves_current(document_store, monkeypatch):
    store, source, _ = document_store
    (source / "good.md").write_text("Original content.")
    first = store.build()
    pointer = (store.index_path / "CURRENT.enc").read_bytes()
    write = Crypto.write

    def fail_pointer(self, path, kind, *args, **kwargs):
        if kind == "index_current":
            raise OSError("simulated publication failure")
        return write(self, path, kind, *args, **kwargs)

    monkeypatch.setattr(Crypto, "write", fail_pointer)
    (source / "good.md").write_text("Changed content.")
    with pytest.raises(OSError, match="simulated publication"):
        store.build(refresh=True)
    assert (store.index_path / "CURRENT.enc").read_bytes() == pointer
    assert store.pin() == first["snapshot_id"]
    assert store.search("Original").chunks[0].text == "Original content."


def test_index_capacity_failure_preserves_current(document_store, monkeypatch):
    from loom.documents import index as documents

    store, source, _ = document_store
    (source / "good.md").write_text("Original content.")
    first = store.build()
    _, image = Crypto(store.config).read(
        store.index_path / "snapshots" / first["snapshot_id"] / "index.sqlite.enc",
        "index_sqlite",
        512 * 1024**2,
    )
    monkeypatch.setattr(documents, "MAXIMUM", len(image) + 4096)
    (source / "large.md").write_text("\n\n".join(f"Added paragraph {i}." for i in range(100)))
    with pytest.raises(ToolFailure) as failure:
        store.build(refresh=True)
    assert failure.value.code == "index_capacity_exceeded"
    assert store.pin() == first["snapshot_id"]


def test_pdf_stream_limit_precedes_extraction(monkeypatch):
    from loom.documents.loaders import pdf

    class Contents:
        def get_data(self):
            return b"x" * (16 * 1024**2 + 1)

    class Page:
        def get_contents(self):
            return Contents()

        def extract_text(self):
            pytest.fail("oversized stream reached text extraction")

    class Reader:
        is_encrypted = False
        pages = [Page()]

    monkeypatch.setattr("pypdf.PdfReader", lambda source: Reader())
    with pytest.raises(ToolFailure) as failure:
        list(pdf(b"placeholder", []))
    assert failure.value.code == "pdf_stream_limit"


def test_docx_malformed_xml_and_partial_content(document_store):
    from io import BytesIO
    from zipfile import ZipFile

    store, source, _ = document_store
    document = Document()
    document.add_paragraph("Searchable body.")
    document.sections[0].header.paragraphs[0].text = "Excluded header."
    document.save(str(source / "partial.docx"))
    data = (source / "partial.docx").read_bytes()
    with ZipFile(BytesIO(data)) as original, ZipFile(source / "broken.docx", "w") as broken:
        for info in original.infolist():
            broken.writestr(
                info,
                b"<broken" if info.filename == "word/document.xml" else original.read(info),
            )
    report = store.build()
    assert report["indexed"] == 1
    assert report["partially_extracted"] == 1
    assert report["excluded"] == 1
    assert report["files"][0]["reason"] == "docx_parse_failed"
    assert snapshot_rows(store, report["snapshot_id"])[0][2] == "Searchable body."


def test_query_clipping_is_explicit_and_bounded(document_store):
    store, source, _ = document_store
    (source / "chunks.md").write_text("\n\n".join("word " * 150 for _ in range(8)))
    store.build()
    result = store.search("word")
    assert result.clipped
    assert 0 < len(result.chunks) < 5
    assert len(canonical(result)) <= 4096
    assert [item.code for item in result.diagnostics] == ["retrieval_context_clipped"]


def test_ingestion_warning_reaches_agent_result(document_store):
    store, source, _ = document_store
    (source / "good.md").write_text("Some indexed content.")
    (source / "excluded.txt").write_text("Not indexed.")
    store.build()
    model = ScriptedModel(
        [
            ModelResponse(
                finish_reason="tool_calls",
                tool_calls=(
                    ToolCall(
                        call_id="search",
                        provider_call_id="search",
                        name="search_docs",
                        arguments={"query": "unrelated"},
                    ),
                ),
            ),
            ModelResponse(
                output={"status": "abstained", "answer": "Insufficient evidence.", "sources": []}
            ),
        ]
    )
    result = Agent(
        name="reader",
        instructions="Answer from indexed documents.",
        data=[store],
        model=model,
        config=store.config,
        output_schema="document_qa",
    ).run(input="A question")
    assert result.status == "completed"
    assert any(item.code == "index_coverage_incomplete" for item in result.diagnostics)


@pytest.fixture
def query_cache():
    from loom.documents import store as documents

    def clear():
        with documents._cache_guard:
            assert not documents._loads
            assert all(entry.leases == 0 for entry in documents._cache.values())
            assert documents._charged == sum(
                2 * entry.size + 64 * documents.MIB for entry in documents._cache.values()
            )
            for entry in documents._cache.values():
                entry.database.close()
            documents._cache.clear()
            documents._charged = 0

    clear()
    yield documents
    clear()


async def wait_until(predicate):
    def wait():
        deadline = time.monotonic() + 5
        while not predicate():
            if time.monotonic() >= deadline:
                raise TimeoutError("Worker did not reach the expected state")
            time.sleep(0.005)

    await asyncio.to_thread(wait)


@pytest.mark.asyncio
async def test_single_flight_survives_one_cancelled_waiter(
    document_store, query_cache, monkeypatch
):
    store, source, _ = document_store
    (source / "good.md").write_text("Shared content.")
    snapshot = store.build()["snapshot_id"]
    started, release = threading.Event(), threading.Event()
    original = Crypto.read
    reads = []

    def delayed_read(self, path, kind, maximum):
        if kind == "index_sqlite":
            reads.append(path)
            started.set()
            assert release.wait(5)
        return original(self, path, kind, maximum)

    monkeypatch.setattr(Crypto, "read", delayed_read)
    first = asyncio.create_task(store.asearch("Shared", snapshot_id=snapshot))
    second = None
    try:
        await wait_until(started.is_set)
        second = asyncio.create_task(store.asearch("content", snapshot_id=snapshot))
        await wait_until(lambda: any(load.waiters == 2 for load in query_cache._loads.values()))
        first.cancel()
        with pytest.raises(asyncio.CancelledError):
            await first
        size = (store.index_path / "snapshots" / snapshot / "index.sqlite.enc").stat().st_size
        assert query_cache._charged == 4 * size + 128 * query_cache.MIB
        release.set()
        result = await second
        assert result.snapshot_id == snapshot
        assert result.chunks[0].text == "Shared content."
        assert len(reads) == 1
        assert query_cache._charged == 2 * size + 64 * query_cache.MIB
        assert next(iter(query_cache._cache.values())).leases == 0
    finally:
        release.set()
        await asyncio.gather(*(task for task in (first, second) if task), return_exceptions=True)
        await wait_until(lambda: not query_cache._loads)


@pytest.mark.asyncio
async def test_cancelled_query_retains_lease_until_inference_finishes(
    document_store, query_cache, monkeypatch
):
    store, source, _ = document_store
    (source / "good.md").write_text("Content.")
    store.build()
    started, release = threading.Event(), threading.Event()
    original = store._search

    def slow_search(*args):
        started.set()
        assert release.wait(5)
        return original(*args)

    monkeypatch.setattr(store, "_search", slow_search)
    task = asyncio.create_task(store.asearch("content"))
    try:
        await wait_until(started.is_set)
        entry = next(iter(query_cache._cache.values()))
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert entry.leases == 1
        assert query_cache._charged == 2 * entry.size + 128 * query_cache.MIB
    finally:
        release.set()
        await asyncio.gather(task, return_exceptions=True)
        await wait_until(lambda: all(entry.leases == 0 for entry in query_cache._cache.values()))
    assert query_cache._charged == 2 * entry.size + 64 * query_cache.MIB


@pytest.mark.asyncio
async def test_two_snapshot_limit_and_idle_lru_eviction(document_store, query_cache):
    store, source, _ = document_store
    snapshots = []
    for index in range(3):
        (source / "good.md").write_text(f"Generation {index}.")
        snapshots.append(store.build(refresh=index > 0)["snapshot_id"])
    entries = []
    try:
        entries.append(await store._snapshot(snapshots[0], time.monotonic() + 5))
        entries.append(await store._snapshot(snapshots[1], time.monotonic() + 5))
        with pytest.raises(ToolFailure) as failure:
            await store._snapshot(snapshots[2], time.monotonic() + 5)
        assert failure.value.code == "query_memory_exhausted"
        assert failure.value.recoverable
        assert failure.value.details["resident_snapshots"] == 2
        with query_cache._cache_guard:
            entries.pop(0).leases -= 1
        entries.append(await store._snapshot(snapshots[2], time.monotonic() + 5))
        assert {key[1] for key in query_cache._cache} == set(snapshots[1:])
        assert store.pin(snapshots[0]) == snapshots[0]
    finally:
        with query_cache._cache_guard:
            for entry in entries:
                entry.leases -= 1


def test_memory_admission_exact_boundary(query_cache, monkeypatch):
    with monkeypatch.context() as patch:
        patch.setattr(query_cache, "_charged", 3 * 1024**3 - 64 * query_cache.MIB)
        query_cache._space(64 * query_cache.MIB, protect=("index", "snapshot"))
        with pytest.raises(ToolFailure) as failure:
            query_cache._space(64 * query_cache.MIB + 1, protect=("index", "snapshot"))
        assert failure.value.code == "query_memory_exhausted"
        assert failure.value.details["available_bytes"] == 64 * query_cache.MIB


@pytest.mark.skipif(not hasattr(os, "fork"), reason="fork is not available on this platform")
def test_idle_fork_reloads_snapshot_and_executor_state(document_store, query_cache):
    import multiprocessing

    store, source, _ = document_store
    (source / "good.md").write_text("Survives an idle fork.")
    snapshot = store.build()["snapshot_id"]
    assert store.search("idle").snapshot_id == snapshot
    parent, child = multiprocessing.get_context("fork").Pipe()

    def search_after_fork():
        try:
            assert not query_cache._cache and not query_cache._loads
            assert query_cache._charged == 0
            result = store.search("idle", snapshot_id=snapshot)
            child.send((result.snapshot_id, result.chunks[0].text))
        finally:
            child.close()

    process = multiprocessing.get_context("fork").Process(target=search_after_fork)
    try:
        process.start()
        child.close()
        assert parent.poll(15), "forked search did not finish"
        assert parent.recv() == (snapshot, "Survives an idle fork.")
        process.join(5)
        assert process.exitcode == 0
    finally:
        if process.is_alive():
            process.kill()
            process.join()
        parent.close()
